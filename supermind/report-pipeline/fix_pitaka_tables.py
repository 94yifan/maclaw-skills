# -*- coding: utf-8 -*-
"""就地修复 docx 中残留的 Markdown 管道表格段落：转成真表格。
用法：python3 fix_pitaka_tables.py <docx路径>
不重建全文，只替换管道段落块，其余内容逐字节保留。
"""
import sys, shutil, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'PingFang SC'
HEADER_BG = '1F3864'
ZEBRA_BG = 'EDF1F8'
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_eastasia(run_el, font):
    rPr = run_el.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)


def set_cell(cell, text, bold=False, size=Pt(9), color=None, bg=None):
    cell.text = ''
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.15
    r = para.add_run(text)
    r.font.name = FONT
    r.font.size = size
    r.bold = bold
    set_eastasia(r._element, FONT)
    if color is not None:
        r.font.color.rgb = color
    if bg:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), bg)
        cell._tc.get_or_add_tcPr().append(shd)


def build_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=Pt(9.5),
                 color=WHITE, bg=HEADER_BG)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, val in enumerate(row):
            bg = ZEBRA_BG if r_idx % 2 == 1 else None
            set_cell(cells[c_idx], str(val), size=Pt(9), bg=bg)
    return table


def main(path):
    p = Path(path).resolve()
    bak = p.with_suffix('.before-tablefix.docx')
    if not bak.exists():
        shutil.copy2(p, bak)
        print('backup:', bak)
    doc = Document(str(p))

    paras = doc.paragraphs
    blocks, cur = [], []
    for idx, para in enumerate(paras):
        if para.text.strip().startswith('|'):
            cur.append(idx)
        else:
            if cur:
                blocks.append(cur)
                cur = []
    if cur:
        blocks.append(cur)
    print('pipe blocks:', len(blocks))

    for blk in blocks:
        lines = [paras[i].text.strip() for i in blk]
        # 表头 + 分隔行 + 数据行
        if len(lines) < 3:
            print('skip short block', blk[0], len(lines))
            continue
        headers = [c.strip() for c in lines[0].split('|')[1:-1]]
        rows = []
        for ln in lines[2:]:
            cells = [c.strip() for c in ln.split('|')[1:-1]]
            if len(cells) == len(headers):
                rows.append(cells)
            else:
                print('skip mismatched row', len(cells), 'vs', len(headers), ln[:60])
        if not rows:
            print('no rows parsed, skip', blk[0])
            continue

        table = build_table(doc, headers, rows)
        anchor = paras[blk[0]]._p
        anchor.addprevious(table._tbl)          # 表格移到原位置
        for i in blk:
            paras[i]._p.getparent().remove(paras[i]._p)  # 删除原管道段落
        # 表格后补一个空段，避免与下文粘连
        p_el = OxmlElement('w:p')
        table._tbl.addnext(p_el)
        print('repaired block at para', blk[0], '->', len(headers), 'cols', len(rows), 'rows')

    doc.save(str(p))
    print('saved:', p)

    # 校验
    d2 = Document(str(p))
    left = [i for i, x in enumerate(d2.paragraphs) if x.text.strip().startswith('|')]
    print('leftover pipe paragraphs:', len(left))
    print('tables:', len(d2.tables))


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else
         'output/reports/PITAKA-高端手机壳-V1-20261017.docx')
