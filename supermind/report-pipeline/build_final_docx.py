# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0,'.')
from config import ReportSchema, ProjectConfig
from steps.docx_builder import build_chapter_map
from typeset import ReportConfig, ReportTypesetter, DataTable
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

pc = ProjectConfig('project_config_王小卤.json')
cm = build_chapter_map(pc)

charts = [
 ("图1：天猫旗舰店爆款销售对比", "output/charts/chart_brand_comparison_1.png", "数据来源：天猫旗舰店页面实测，2026年10月"),
 ("图2：京东自营爆款销售对比", "output/charts/chart_brand_comparison_2.png", "数据来源：京东自营旗舰店页面实测，2026年10月"),
 ("图3：各品牌核心产品单件价对比", "output/charts/chart_brand_comparison_3.png", "数据来源：天猫与京东旗舰店定价，2026年10月"),
 ("图4：各品牌回头客复购率对比", "output/charts/chart_brand_comparison_4.png", "数据来源：天猫旗舰店回头客标签，2026年10月"),
]

parts = []
parts.append("# 王小卤品牌与赛道分析报告")
parts.append(cm['brand_overview'])
parts.append(cm['product_matrix'])
parts.append(cm['channel_supply_chain'])
parts.append(cm['ch1'])
parts.append(cm['ch2'])
parts.append(cm['ch3_deep_brands'])
parts.append(cm['ch3_summary'])
parts.append(cm['ch3_patterns'])
parts.append(cm['ch3_content_types'])
# chart section (captions; images inserted in post-process)
chart_block = ["## 数据可视化：核心竞品对比", ""]
for cap,_,src in charts:
    chart_block.append(cap)
    chart_block.append(src)
parts.append("\n".join(chart_block))
parts.append(cm['ch4'])
parts.append(cm['ch5'])
parts.append(cm['ch6'])
parts.append(cm['appendix_innovation'])
parts.append(cm['appendix_founder'])

md = "\n\n".join(p for p in parts if p)
open('output/reports/_combined_王小卤.md','w',encoding='utf-8').write(md)
print('combined md chars:', len(md))

cfg = ReportConfig(
    company='燃创咨询（BreaC）',
    title='王小卤品牌与赛道分析报告',
    subtitle='休闲卤味零食赛道',
    client='王小卤',
    framework='品牌概览 → 行业格局 → 竞品五维扫描 → 本品诊断 → 差距对比 → 策略建议 → 创品与创始人研究',
    date='2026年10月12日',
    font='PingFang SC',
)
class TS(ReportTypesetter):
    def _parse_md_table(self, doc, lines, i):
        header_cells = [c.strip() for c in lines[i].split('|')[1:-1]]
        i += 2
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].split('|')[1:-1]]
            if len(cells) == len(header_cells):
                rows.append(cells)
            i += 1
        if header_cells and rows:
            dt = DataTable(title='', headers=header_cells, rows=rows, header_bg='1F3864', zebra_bg='EDF1F8')
            self._add_data_table(doc, dt)
        return i

ts = TS(cfg)
out = 'output/reports/王小卤-休闲卤味零食-V1-20261012.docx'
ts.convert('output/reports/_combined_王小卤.md', out)

# ---- post-process: insert chart images after the 数据可视化 heading ----
def insert_image_after(para, img, width_in=6.0):
    doc = para._p.getparent()
    new_p = para._p.addnext(OxmlElement('w:p'))
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, para._parent)
    np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = np.add_run()
    run.add_picture(img, width=Inches(width_in))
    return np

doc = Document(out)
# find heading para of 数据可视化
target = None
for p in doc.paragraphs:
    if '数据可视化' in p.text and p.style.name.startswith('Heading'):
        target = p
        break
assert target is not None, 'chart heading not found'

from docx.text.paragraph import Paragraph
cur = target._p
for cap,img,src in charts:
    # caption paragraph
    cap_p = OxmlElement('w:p'); cur.addnext(cap_p); cur = cap_p
    cp = Paragraph(cap_p, target._parent); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(cap); r.bold=True; r.font.size=Pt(10.5); r.font.name='PingFang SC'
    r._element.get_or_add_rPr()
    # image paragraph
    img_p = OxmlElement('w:p'); cap_p.addnext(img_p); cur = img_p
    ip = Paragraph(img_p, target._parent); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(img, width=Inches(6.0))
    # source paragraph
    src_p = OxmlElement('w:p'); img_p.addnext(src_p); cur = src_p
    sp = Paragraph(src_p, target._parent); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sp.add_run(src); sr.font.size=Pt(8); sr.font.color.rgb=RGBColor(128,128,128); sr.font.name='PingFang SC'

doc.save(out)
print('saved', out)
