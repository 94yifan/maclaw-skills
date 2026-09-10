#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""春光食品报告 排版后处理：封面页 + 中文字体(eastAsia) + 页码页脚 + 标题分级配色。

在 pipeline 生成的 docx 基础上做增量排版：重建封面页、不动正文/表格/图表。
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x33, 0x55)
GOLD = RGBColor(0xB0, 0x8A, 0x3E)
GRAY = RGBColor(0x66, 0x66, 0x66)
FONT = 'PingFang SC'


def set_run_font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(k), name)


def set_style_font(style, name=FONT):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(k), name)


def fmt(p, size, bold=False, color=None, align=None):
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        if color is not None:
            r.font.color.rgb = color
        set_run_font(r)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

    def add_text(t):
        r = p.add_run(t)
        set_run_font(r); r.font.size = Pt(9); r.font.color.rgb = GRAY

    def add_field(instr):
        f = OxmlElement('w:fldSimple')
        f.set(qn('w:instr'), instr)
        run = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
            rf.set(qn(k), FONT)
        rpr.append(rf)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); rpr.append(sz)
        run.append(rpr)
        t = OxmlElement('w:t'); t.text = '1'; run.append(t)
        f.append(run)
        p._p.append(f)

    add_text('第 ')
    add_field(' PAGE \\* MERGEFORMAT ')
    add_text(' 页 / 共 ')
    add_field(' NUMPAGES \\* MERGEFORMAT ')
    add_text(' 页')


def main(path):
    doc = Document(path)

    # 1) 全局中文字体
    for sname in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4',
                  'Title', 'List Bullet', 'List Number', 'No Spacing'):
        try:
            set_style_font(doc.styles[sname])
        except KeyError:
            pass
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.5

    try:
        doc.styles['Heading 1'].font.color.rgb = NAVY
        doc.styles['Heading 1'].font.size = Pt(18)
        doc.styles['Heading 1'].font.bold = True
        doc.styles['Heading 2'].font.color.rgb = NAVY
        doc.styles['Heading 2'].font.size = Pt(14)
        doc.styles['Heading 2'].font.bold = True
        doc.styles['Heading 3'].font.color.rgb = GOLD
        doc.styles['Heading 3'].font.size = Pt(12)
        doc.styles['Heading 3'].font.bold = True
    except Exception:
        pass

    # 2) 定位封面原始行 + 目录锚点
    title, sub, ind, date = '', '', '', ''
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if '品牌扫描报告' in t and not title:
            title = t
        elif t == '深度品牌扫描' and not sub:
            sub = t
        elif '椰基食品饮料 / 海南旅游' in t and not ind:
            ind = t
        elif re.match(r'^\d{4}年\d{2}月\d{2}日$', t) and not date:
            date = t

    toc_p = None
    for p in doc.paragraphs:
        if p.text.strip() == '目录':
            toc_p = p
            break
    if toc_p is None:
        raise SystemExit('未找到目录段，无法定位封面')

    # 删除目录之前的所有段落
    for p in list(doc.paragraphs):
        if p._p is toc_p._p:
            break
        p._p.getparent().remove(p._p)

    # 3) 重建封面页
    cover = [
        ('燃创咨询（BreaC）', 12, True, GOLD),
        ('', 10, False, None),
        ('', 10, False, None),
        (title, 22, True, NAVY),
        (sub, 15, False, NAVY),
        (ind, 11, False, GRAY),
        ('', 10, False, None),
        ('委托方：春光食品（海南春光食品有限公司）', 11, False, GRAY),
        ('研究框架：五维扫描 / 产业链卡口 / 内容五分类 / 机会地图 / 创品策略', 11, False, GRAY),
        (date, 11, False, GRAY),
    ]
    for text, size, bold, color in cover:
        np = toc_p.insert_paragraph_before(text)
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
        np.paragraph_format.space_after = Pt(8)
        fmt(np, size, bold=bold, color=color)

    # 目录另起一页
    toc_p.paragraph_format.page_break_before = True

    # 4) 页脚页码 + 页边距
    add_page_number_footer(doc.sections[0])
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.6)
        s.right_margin = Cm(2.6)

    doc.save(path)
    print('排版完成:', path)


if __name__ == '__main__':
    main(sys.argv[1])
